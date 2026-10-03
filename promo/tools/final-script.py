# 최종 발표 대본 (5분, 홍보 영상을 맨 처음에 튼다) — 2026-10-03 동근님 최종 PPT(9장) 기준.
#
#   python3 promo/tools/final-script.py <QR 넣은 최종 pptx> <출력 폴더>
#
# 슬라이드 노트에 장마다 시간·대사를 넣은 pptx를 다시 저장하고, 대본 .md / .html(→ Chrome으로 PDF)을 만든다.
# 말 빠르기는 1초에 4~4.5음절로 잡았다. 합계 4:45 (15초 여유).
import re
import sys
from pathlib import Path

from pptx import Presentation

SRC, OUT = Path(sys.argv[1]), Path(sys.argv[2])
OUT.mkdir(parents=True, exist_ok=True)

VIDEO = ('0:00–0:50', '홍보 영상 (50초)', '[발표 시작 전에 영상을 먼저 튼다. 말하지 않고 기다린다. 영상이 끝나면 1장을 띄운다]')
SLIDES = [
    ('0:50–1:05', '표지',
     '안녕하세요, 2gether 팀의 김다미입니다. 방금 보신 서비스, 어울림입니다. '
     '한국에 온 외국인 유학생에게 「우리 조」를 만들어 주는 캠퍼스 커뮤니티입니다.'),
    ('1:05–1:40', '정착 문제',
     '학령인구가 줄면서 대학은 유학생으로 빈자리를 채우고 있습니다. 그런데 데려오는 것만큼 중요한 게 정착입니다. '
     '교환학생은 SG MAPLE 버디로 잘 적응하지만, 글로벌 버디는 신입생만 대상이라 편입생이나 학위과정 유학생은 기댈 곳이 없습니다. '
     '1:1은 부담스럽고, 일회성 모임은 금방 끝나죠. 그래서 저희는 한 학기 내내 이어지는 그룹을 만들었습니다.'),
    ('1:40–1:45', '50초로 보는 어울림',
     '방금 영상에서 보신 기능을 하나씩 짚어 보겠습니다. [바로 다음 장으로 — 영상은 이미 틀었다]'),
    ('1:45–2:20', '우리 조',
     '핵심은 「우리 조」입니다. 학교와 나라, 언어만 고르면 같은 나라 친구 3~4명과 한국인 멘토 1명이 있는 조에 자동으로 들어갑니다. '
     '그룹은 무료로 운영하고, 수익은 이벤트에서 냅니다. 멘토는 조의 관리자로 게시판과 전체 활동을 이끌고, '
     '매주 그룹 미션으로 처음의 어색함을 같이 깹니다.'),
    ('2:20–3:05', '한국 학생은 멘토로',
     '한국 학생은 멘토로 들어옵니다. 돈을 받고 조를 관리하기 때문에, 어울림 운영팀이 학생증을 직접 확인한 뒤에만 들어올 수 있습니다. '
     '외국인 조는 계속 유지되고, 멘토는 한 학기씩 맡아 연장하거나 다음 멘토에게 넘깁니다. '
     '보상은 활동증명서와 비교과 인정, 수익의 30% 배분, 이벤트 무료 참가입니다. '
     '관심이 적은 나라는 운영팀이 직접 맡고, 학생이 세 명이 안 되는 나라는 7일 뒤 같은 지역 조로 묶습니다.'),
    ('3:05–3:30', '커뮤니티',
     '그룹 밖의 가벼운 도움은 커뮤니티에서 해결합니다. 한국인 멘토와 유학생 모두 쓰는 게시판이고, '
     '질문을 쓰면 성대 공식 안내부터 먼저 보여 줍니다. 물건을 빌리거나 오늘 같이 밥 먹을 사람도 찾을 수 있고, 포인트는 선택입니다.'),
    ('3:30–4:00', '이벤트',
     '수익은 이벤트에서 냅니다. 사주 카페, 한복 고궁 산책 같은 한국 문화 체험과 '
     '필리핀 Undas, 베트남 여성의 날 같은 나라별 명절 이벤트를 엽니다. 참가비의 15%가 어울림 수수료입니다. '
     '이벤트 사진을 SNS에 올려 홍보해 주면 포인트를 드립니다.'),
    ('4:00–4:25', '수익 구조',
     '정리하면, 우리 조는 무료, 이벤트 수수료가 메인, 포인트가 보조 수익입니다. 그중 30%는 한국인 멘토가 실적대로 나눠 받습니다. '
     '예를 들어 만 이천 원짜리 Undas 저녁이면 어울림 몫은 천팔백 원입니다.'),
    ('4:25–4:45', '마무리',
     '다음 단계로는 성대 국제처와 함께, 새로 온 유학생이 첫날 QR 하나로 우리 조를 만나게 하려고 합니다. '
     '한국에서의 첫날부터, 우리 조와 함께. 어울림이었습니다. 감사합니다.'),
]

prs = Presentation(SRC)
if len(prs.slides) != len(SLIDES):
    raise SystemExit(f'장 수가 다르다: PPT {len(prs.slides)}장 / 대본 {len(SLIDES)}장')
for s, (tm, name, said) in zip(prs.slides, SLIDES):
    s.notes_slide.notes_text_frame.text = f'[{tm}] {name}\n\n{said}'
prs.save(SRC)
print('노트 갱신:', SRC)

syl = lambda t: len(re.findall(r'[가-힣0-9A-Za-z]', re.sub(r'\[.*?\]', '', t)))
def secs(tm):
    a, b = [int(x.split(':')[0]) * 60 + int(x.split(':')[1]) for x in tm.split('–')]
    return b - a
for tm, name, said in SLIDES:
    d = secs(tm)
    print(f'{name:14} {d:3}초 {syl(said):4}음절 {syl(said) / d if d else 0:.1f}/초')

rows = [(VIDEO[0], VIDEO[1], VIDEO[2], 0)] + [(tm, name, said, i) for i, (tm, name, said) in enumerate(SLIDES, 1)]
md = ['# 어울림 발표 대본 — 5분 (최종, 2026-10-03)', '',
      '발표 김다미 · 팀 2gether · `어울림_발표_최종.pptx` (PDF판 `어울림_발표_최종.pdf`) · **영상 50초를 맨 처음에 틀고 시작** · 합계 4분 45초 (15초 여유)', '',
      '| 순서 | 시간 | 장 |', '|---|---|---|']
md += [f'| {"영상" if i == 0 else f"{i}장"} | {tm} | {name} |' for tm, name, _, i in rows] + ['']
for tm, name, said, i in rows:
    md += [f'## {"영상" if i == 0 else f"{i}장"} · {name}  `{tm}`', '', said, '']
md += ['---', '', '- 영상 파일: `어울림_홍보영상_50초_최종.mp4` — 발표 PC에서 미리 한 번 틀어 보기 (소리 있음)',
       '- 모든 장 오른쪽 아래에 앱 QR (eoullim-skku.vercel.app). 마지막 장은 큰 QR',
       '- 시간이 밀리면 6장(커뮤니티)을 한 문장으로: 「가벼운 도움은 커뮤니티에서, 공식 안내부터 먼저 보여 줍니다.」']
(OUT / '어울림_발표대본_최종.md').write_text('\n'.join(md) + '\n', encoding='utf-8')

table = ''.join(f'<tr><td>{"영상" if i == 0 else f"{i}장"}</td><td>{tm}</td><td>{name}</td></tr>' for tm, name, _, i in rows)
body = ''.join(
    f'<section class="{"video" if i == 0 else ""}"><h2><span>{"▶" if i == 0 else i}</span>{name}<time>{tm}</time></h2><p>{said}</p></section>'
    for tm, name, said, i in rows)
html = f'''<!doctype html><html lang="ko"><head><meta charset="utf-8"><title>어울림 발표 대본 최종</title><style>
@page {{ size: A4; margin: 15mm 16mm; }} body {{ font-family: 'NanumGothic', sans-serif; color: #0F1B2E; font-size: 11pt; line-height: 1.7; word-break: keep-all; margin: 0; }}
h1 {{ font-size: 19pt; margin: 0 0 2px; }} .sub {{ color: #5D6577; margin: 0 0 12px; font-size: 10pt; }}
table {{ border-collapse: collapse; width: 100%; margin-bottom: 14px; font-size: 9.5pt; }} td {{ padding: 3px 8px; border-bottom: 1px solid #E4E6EB; }}
tr td:first-child {{ width: 40px; color: #5B4DFF; font-weight: 700; }} tr td:nth-child(2) {{ width: 100px; font-family: Menlo, monospace; }}
section {{ break-inside: avoid; margin: 0 0 11px; }} section.video p {{ color: #5D6577; }}
h2 {{ font-size: 12.5pt; margin: 0 0 3px; display: flex; gap: 10px; align-items: center; }}
h2 span {{ display: inline-grid; place-items: center; width: 26px; height: 26px; border-radius: 50%; background: #5B4DFF; color: #fff; font-size: 10.5pt; }}
h2 time {{ margin-left: auto; font-family: Menlo, monospace; font-size: 10pt; color: #2B1FB8; background: #EEECFF; padding: 2px 8px; border-radius: 6px; }}
p {{ margin: 0; padding-left: 36px; }} .tip {{ margin-top: 12px; padding: 9px 12px; background: #F5F6F8; border-radius: 8px; font-size: 9.5pt; }}
</style></head><body><h1>어울림 발표 대본 — 5분 (최종)</h1>
<p class="sub">발표 김다미 · 팀 2gether · 영상 50초를 맨 처음에 틀고 시작 · 합계 4분 45초 (15초 여유) · 2026-10-03</p>
<table>{table}</table>{body}
<div class="tip">· 영상 파일 「어울림_홍보영상_50초_최종.mp4」 — 발표 PC에서 미리 한 번 틀어 보기 (소리 있음)<br>· 모든 장 오른쪽 아래에 앱 QR, 마지막 장은 큰 QR<br>· 시간이 밀리면 6장(커뮤니티)을 한 문장으로: 「가벼운 도움은 커뮤니티에서, 공식 안내부터 먼저 보여 줍니다.」</div>
</body></html>'''
(OUT / '어울림_발표대본_최종.html').write_text(html, encoding='utf-8')
print('대본:', OUT / '어울림_발표대본_최종.md')
