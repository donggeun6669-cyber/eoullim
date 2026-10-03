# 발표 PPT의 모든 장 오른쪽 아래에 앱 QR을 넣는다 (이미 큰 QR이 있는 장은 건너뜀).
#   python3 promo/tools/add-qr.py <입력 pptx> <출력 pptx>
# 동근님이 고친 최종본은 그대로 두고 새 파일로 저장한다. QR = https://eoullim-skku.vercel.app
import sys
from pathlib import Path

import qrcode
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches

ROOT = Path(__file__).resolve().parent.parent
SRC, OUT = sys.argv[1], sys.argv[2]
URL = 'https://eoullim-skku.vercel.app'
png = ROOT / 'out/qr-small.png'
png.parent.mkdir(parents=True, exist_ok=True)
q = qrcode.QRCode(error_correction=qrcode.constants.ERROR_CORRECT_M, border=1, box_size=20)
q.add_data(URL)
q.make(fit=True)
q.make_image(fill_color='#0F1B2E', back_color='white').save(png)

prs = Presentation(SRC)
W, H = prs.slide_width, prs.slide_height
SIZE, PAD = Inches(0.68), Inches(0.05)  # 휴대폰 사진 아래 빈 공간(6.6인치 아래)에 들어가는 크기
added = []
for i, s in enumerate(prs.slides, 1):
    # 이미 QR이 크게 들어간 장(마무리)은 건너뛴다: 정사각형 그림이 2인치 이상
    if any(sh.shape_type == 13 and abs(sh.width - sh.height) < Inches(0.05) and sh.width > Inches(2) for sh in s.shapes):
        continue
    x, y = W - SIZE - Inches(0.1), H - SIZE - Inches(0.11)
    bg = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x - PAD, y - PAD, SIZE + PAD * 2, SIZE + PAD * 2)
    bg.fill.solid()
    bg.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    bg.line.color.rgb = RGBColor(0xE4, 0xE6, 0xEB)
    bg.shadow.inherit = False
    bg.adjustments[0] = 0.12
    bg.name = 'QR background'
    pic = s.shapes.add_picture(str(png), x, y, SIZE, SIZE)
    pic.name = 'QR eoullim-skku.vercel.app'
    added.append(i)
prs.save(OUT)
print('QR 넣은 장:', added, '→', OUT)
