# 최종 발표판의 앱 화면 사진을 새로 찍은 화면으로 바꾼다 (글·QR·노트는 그대로).
#   python3 promo/tools/refresh-final-shots.py <최종 pptx>
# 2026-10-03: 첫 화면 「가입 없음」 문구 삭제 · 조당 멘토 1명(필리핀 2조 멘토 = 서연)에 맞춘 그룹방·채팅
import sys
from pathlib import Path

from PIL import Image, ImageDraw, ImageFilter
from pptx import Presentation
from pptx.util import Inches

SRC = sys.argv[1]
ROOT = Path(__file__).resolve().parent.parent
SHOTS = ROOT / 'assets/screens'
FRAMES = ROOT / 'out/frames'
FRAMES.mkdir(parents=True, exist_ok=True)


def framed(name):
    """build-deck.py 와 같은 폰 테두리 (남색 테두리 + 그림자)"""
    out = FRAMES / f'{name}.png'
    im = Image.open(SHOTS / f'{name}.png').convert('RGB')
    s = 640 / im.width
    im = im.resize((640, round(im.height * s)), Image.LANCZOS)
    bez, r_out, pad = 16, 64, 40
    W, H = im.width + bez * 2, im.height + bez * 2
    canvas = Image.new('RGBA', (W + pad * 2, H + pad * 2), (0, 0, 0, 0))
    sh = Image.new('L', canvas.size, 0)
    ImageDraw.Draw(sh).rounded_rectangle((pad, pad + 14, pad + W, pad + H + 14), r_out, fill=70)
    sh = sh.filter(ImageFilter.GaussianBlur(18))
    canvas.paste(Image.new('RGBA', canvas.size, (15, 27, 46, 255)), (0, 0), sh)
    body = Image.new('RGBA', (W, H), (0, 0, 0, 0))
    m = Image.new('L', (W, H), 0)
    ImageDraw.Draw(m).rounded_rectangle((0, 0, W - 1, H - 1), r_out, fill=255)
    body.paste(Image.new('RGBA', (W, H), (15, 27, 46, 255)), (0, 0), m)
    m2 = Image.new('L', im.size, 0)
    ImageDraw.Draw(m2).rounded_rectangle((0, 0, im.width - 1, im.height - 1), r_out - bez, fill=255)
    body.paste(im, (bez, bez), m2)
    canvas.alpha_composite(body, (pad, pad))
    canvas.save(out)
    return out




def pics_at(slide, left_in, top_in):
    return [sh for sh in slide.shapes if sh.shape_type == 13 and abs(sh.left - Inches(left_in)) < Inches(0.05) and abs(sh.top - Inches(top_in)) < Inches(0.05)]


def swap(slide, left_in, top_in, name):
    olds = pics_at(slide, left_in, top_in)
    if not olds:
        raise SystemExit(f'그림 없음: ({left_in}, {top_in})')
    top = olds[-1]  # 맨 위에 보이는 그림 자리에 넣고, 겹친 옛 그림은 모두 지운다
    new = slide.shapes.add_picture(str(framed(name)), top.left, top.top, height=top.height)
    top._element.addprevious(new._element)
    for o in olds:
        o._element.getparent().remove(o._element)


prs = Presentation(SRC)
S = list(prs.slides)
swap(S[0], 8.55, 0.42, 'a01_welcome')     # 1장 표지 폰
swap(S[3], 10.25, 0.75, 'a03_group')      # 4장 오른쪽 그룹방
swap(S[4], 10.25, 0.75, 'b03_group_chat_kr')  # 5장 오른쪽 채팅
# 3장 「50초로 보는 어울림」: 새 영상의 16.5초 장면 (우리 조 멘토 = 서연)
import subprocess
video = ROOT / 'out/어울림_홍보_50초_1080p_사운드.mp4'
still = FRAMES / 'video_still_16_5.png'
subprocess.run(['ffmpeg', '-v', 'error', '-y', '-ss', '16.5', '-i', str(video), '-frames:v', '1', str(still)], check=True)
olds = pics_at(S[2], 0.95, 1.0)
if olds:
    o = olds[-1]
    new = S[2].shapes.add_picture(str(still), o.left, o.top, o.width, o.height)
    o._element.addprevious(new._element)
    for x in olds:
        x._element.getparent().remove(x._element)
prs.save(SRC)
print('사진 교체:', SRC)
