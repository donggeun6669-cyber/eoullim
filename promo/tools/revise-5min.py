# 5분 발표판 수정본 만들기 (2026-10-03).
#
#   python3 promo/tools/revise-5min.py <동근님이 고친 pptx> <출력 폴더>
#
# 동근님이 드라이브(구글 슬라이드)에서 고친 판을 그대로 두고, 그 위에서 바꾼 새 파일을 만든다.
#   - 3장: 구글 슬라이드가 지운 홍보 영상을 다시 넣는다 (같은 자리, 클릭하면 재생)
#   - 4장: 동근님 문구는 그대로, 사진 두 장만 새 화면으로 (그룹방 「빈자리」 없음 · 시작 화면 「한국 친구」 없음)
#   - 5장: 「한국 친구」 → 「인증된 한국인 멘토」 (한 학기 교대 · 보상 · 파견 예정자 · 지역 조)
#   - 6장 아래 한 줄: 「관심 있는 한국 학생」 → 「관심 있는 한국인 멘토」 (앱 표시와 맞춤)
#   - 슬라이드 노트(대본)에서 바뀐 장만 고친다. 대본 PDF용 HTML·MD도 노트에서 만든다.
import subprocess
import sys
from pathlib import Path

from PIL import Image, ImageDraw, ImageFilter
from pptx import Presentation

ROOT = Path(__file__).resolve().parent.parent
SHOTS = ROOT / 'assets/screens'
FRAMES = ROOT / 'out/frames'
FRAMES.mkdir(parents=True, exist_ok=True)
VIDEO = ROOT / 'out/어울림_홍보_50초_1080p_사운드.mp4'
SRC, OUT = Path(sys.argv[1]), Path(sys.argv[2])
OUT.mkdir(parents=True, exist_ok=True)


def framed(name):
    """build-deck.py 와 같은 폰 테두리 (남색 테두리 + 그림자)"""
    out = FRAMES / f'{name}.png'
    im = Image.open(SHOTS / f'{name}.png').convert('RGB')
    s = 640 / im.width
    im = im.resize((640, round(im.height * s)), Image.LANCZOS)
    bez, r_out, pad = 16, 64, 40
    W, H = im.width + bez * 2, im.height + bez * 2
    canvas = Image.new('RGBA', (W + pad * 2, H + pad * 2), (0, 0, 0, 0))
    sh = Image.new('L', canvas.size, 0)
    ImageDraw.Draw(sh).rounded_rectangle((pad, pad + 14, pad + W, pad + H + 14), r_out, fill=70)
    sh = sh.filter(ImageFilter.GaussianBlur(18))
    canvas.paste(Image.new('RGBA', canvas.size, (15, 27, 46, 255)), (0, 0), sh)
    body = Image.new('RGBA', (W, H), (0, 0, 0, 0))
    m = Image.new('L', (W, H), 0)
    ImageDraw.Draw(m).rounded_rectangle((0, 0, W - 1, H - 1), r_out, fill=255)
    body.paste(Image.new('RGBA', (W, H), (15, 27, 46, 255)), (0, 0), m)
    m2 = Image.new('L', im.size, 0)
    ImageDraw.Draw(m2).rounded_rectangle((0, 0, im.width - 1, im.height - 1), r_out - bez, fill=255)
    body.paste(im, (bez, bez), m2)
    canvas.alpha_composite(body, (pad, pad))
    canvas.save(out)
    return out


def shape(slide, name):
    for sh in slide.shapes:
        if sh.name == name:
            return sh
    raise SystemExit(f'모양을 못 찾음: {name}')


def replace_picture(slide, name, img):
    """같은 자리·같은 높이에 새 그림을 넣고, 쌓는 순서도 그대로 둔다"""
    old = shape(slide, name)
    new = slide.shapes.add_picture(str(img), old.left, old.top, height=old.height)
    old._element.addprevious(new._element)
    old._element.getparent().remove(old._element)
    new.name = name


def set_text(slide, name, text):
    """첫 글자 모양(크기·굵기·색)은 그대로 두고 글만 바꾼다"""
    tf = shape(slide, name).text_frame
    p0 = tf.paragraphs[0]
    for p in tf.paragraphs[1:]:
        p._p.getparent().remove(p._p)
    runs = p0.runs
    runs[0].text = text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)


def set_notes(slide, tm, title, said):
    slide.notes_slide.notes_text_frame.text = f'[{tm}] {title}\n\n{said}'


prs = Presentation(SRC)
S = list(prs.slides)

# 3장 — 영상 다시 넣기
s3 = S[2]
pic = shape(s3, 'Google Shape;122;p3')
poster = FRAMES / 'video_poster.png'
subprocess.run(['ffmpeg', '-v', 'error', '-y', '-ss', '16.5', '-i', str(VIDEO), '-frames:v', '1', str(poster)], check=True)
mov = s3.shapes.add_movie(str(VIDEO), pic.left, pic.top, pic.width, pic.height, poster_frame_image=str(poster), mime_type='video/mp4')
pic._element.addprevious(mov._element)
pic._element.getparent().remove(pic._element)
n3 = s3.notes_slide.notes_text_frame.text.replace('「우리 조 → 한국 친구 → 이벤트」', '「우리 조 → 한국인 멘토 → 이벤트」')
s3.notes_slide.notes_text_frame.text = n3

# 4장 — 동근님 문구는 그대로, 사진 두 장만 새 화면으로 (오른쪽 그룹방: 빈자리 없음 / 왼쪽 시작 화면: 「한국 친구」 문구 없음)
replace_picture(S[3], 'Google Shape;145;p4', framed('a02_setup'))
replace_picture(S[3], 'Google Shape;146;p4', framed('a03_group'))

# 5장 — 인증된 한국인 멘토
s5 = S[4]
set_text(s5, 'Google Shape;153;p5', '멘토')
set_text(s5, 'Google Shape;154;p5', '한국 학생은 「인증된 멘토」로')
set_text(s5, 'Google Shape;155;p5', '돈을 받고 조를 관리하니까 — 운영팀이 학생증을 직접 확인하고 1인 1코드. (지금은 데모라 생략 가능)')
set_text(s5, 'Google Shape;158;p5', '한 학기 약속, 멘토는 교대')
set_text(s5, 'Google Shape;159;p5', '조는 학기가 바뀌어도 계속 남고, 멘토는 한 학기씩 맡아 연장하거나 다음 멘토에게 인계합니다.')
set_text(s5, 'Google Shape;162;p5', '보상: 활동 인정 · 수익 배분 · 이벤트 무료')
set_text(s5, 'Google Shape;163;p5', '학기 활동증명서·비교과 인정(국제처 협업) / 어울림 수익의 30% 실적 배분 / 자기 조와 가는 이벤트 무료.')
set_text(s5, 'Google Shape;166;p5', '관심 적은 나라 · 학생이 적은 나라도')
set_text(s5, 'Google Shape;167;p5', '그 나라로 교환학생 갈 한국 학생을 멘토로 먼저 모집. 3명이 안 되는 나라는 7일 뒤 지역 조로 함께.')
set_text(s5, 'Google Shape;168;p5', '오른쪽: 필리핀 교환학생 예정인 멘토 서연의 화면 — Joy의 메시지가 한국어로 (15개 언어 자동 번역)')
replace_picture(s5, 'Google Shape;169;p5', framed('b00_verify'))
replace_picture(s5, 'Google Shape;170;p5', framed('b03_group_chat_kr'))
set_notes(s5, '2:20–2:50', '한국인 멘토',
          '한국 학생은 돈을 받고 조를 관리하니까, 운영팀이 학생증을 확인한 뒤 인증 코드를 줍니다. '
          '멘토는 한 학기씩 맡고 교대해도, 조는 계속 남습니다. 보상은 활동 인정, 수익 배분, 이벤트 무료 참가입니다. '
          '그 나라로 교환학생 갈 학생을 멘토로 먼저 모집하고, 학생이 적은 나라는 지역 조로 묶습니다.')

# 6장 아래 한 줄
set_text(S[5], 'Google Shape;192;p6', '나라별 행사에는 「그 나라에 관심 있는 한국인 멘토 n명」이 보여 참가자가 넓어집니다.')

out = OUT / '어울림_발표_5분_수정본.pptx'
prs.save(out)
print('저장:', out)

# 대본 — 슬라이드 노트에서 그대로 뽑는다 (동근님이 고친 노트도 같이 반영됨)
rows = []
for s in prs.slides:
    t = s.notes_slide.notes_text_frame.text if s.has_notes_slide else ''
    head, _, body = t.partition('\n\n')
    tm, _, name = head.strip().lstrip('[').partition('] ')
    rows.append((tm, name, body.strip()))
md = ['# 어울림 발표 대본 — 5분판 수정본 (2026-10-03)', '',
      '발표: 김다미 · 팀 2gether · 파일 `어울림_발표_5분_수정본.pptx` (PDF판 `어울림_발표_5분_수정본.pdf`) · 합계 **4분 45초**', '',
      '| 장 | 시간 | 내용 |', '|---|---|---|'] + [f'| {i} | {tm} | {n} |' for i, (tm, n, _) in enumerate(rows, 1)] + ['']
for i, (tm, n, said) in enumerate(rows, 1):
    md += [f'## {i}. {n}  `{tm}`', '', said, '']
md += ['---', '', '- 3장 영상은 화면을 클릭하면 재생 (소리 있음). **PowerPoint·Keynote에서 열 것** — 구글 슬라이드에서 고치면 영상이 빠진다',
       '- PDF로 발표하면 영상이 안 나온다 → 영상 파일을 따로 틀고 돌아온다']
(OUT / '어울림_발표대본_5분_수정본.md').write_text('\n'.join(md) + '\n', encoding='utf-8')
body = ''.join(f'<section><h2><span>{i}</span>{n}<time>{tm}</time></h2><p>{said}</p></section>' for i, (tm, n, said) in enumerate(rows, 1))
table = ''.join(f'<tr><td>{i}</td><td>{tm}</td><td>{n}</td></tr>' for i, (tm, n, _) in enumerate(rows, 1))
html = f'''<!doctype html><html lang="ko"><head><meta charset="utf-8"><title>어울림 발표 대본 5분 수정본</title><style>
@page {{ size: A4; margin: 16mm; }} body {{ font-family: 'NanumGothic', sans-serif; color: #0F1B2E; font-size: 11pt; line-height: 1.7; word-break: keep-all; margin: 0; }}
h1 {{ font-size: 20pt; margin: 0 0 4px; }} .sub {{ color: #5D6577; margin: 0 0 14px; }}
table {{ border-collapse: collapse; width: 100%; margin-bottom: 18px; font-size: 10pt; }} td {{ padding: 4px 8px; border-bottom: 1px solid #E4E6EB; }} tr td:first-child {{ width: 30px; color: #5B4DFF; font-weight: 700; }} tr td:nth-child(2) {{ width: 110px; font-family: Menlo, monospace; }}
section {{ break-inside: avoid; margin: 0 0 14px; }} h2 {{ font-size: 13pt; margin: 0 0 4px; display: flex; gap: 10px; align-items: center; }}
h2 span {{ display: inline-grid; place-items: center; width: 26px; height: 26px; border-radius: 50%; background: #5B4DFF; color: #fff; font-size: 11pt; }}
h2 time {{ margin-left: auto; font-family: Menlo, monospace; font-size: 10pt; color: #2B1FB8; background: #EEECFF; padding: 2px 8px; border-radius: 6px; }}
p {{ margin: 0; padding-left: 36px; white-space: pre-line; }} .tip {{ margin-top: 18px; padding: 10px 12px; background: #F5F6F8; border-radius: 8px; font-size: 10pt; }}
</style></head><body><h1>어울림 발표 대본 — 5분판 수정본</h1><p class="sub" style="padding:0">발표 김다미 · 팀 2gether · 홍보 영상 50초 포함 합계 4분 45초 · 2026-10-03</p>
<table>{table}</table>{body}
<div class="tip">· 3장 영상은 화면을 클릭하면 재생(소리 있음). PowerPoint·Keynote에서 열 것 — 구글 슬라이드에서 고치면 영상이 빠진다<br>· PDF로 발표하면 영상이 안 나온다 → 영상 파일을 따로 틀고 돌아온다</div></body></html>'''
(OUT / '어울림_발표대본_5분_수정본.html').write_text(html, encoding='utf-8')
print('대본:', OUT / '어울림_발표대본_5분_수정본.md')
